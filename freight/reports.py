import csv, io, json, zipfile, html
from .core import money

def safe_cell(value):
    s=str('' if value is None else value)
    return "'"+s if s.lstrip().startswith(('=','+','-','@','\t','\r')) else s

def csv_bytes(fields,rows):
    s=io.StringIO(newline='');w=csv.writer(s);w.writerow(fields)
    for row in rows: w.writerow([safe_cell(x) for x in row])
    return ('\ufeff'+s.getvalue()).encode('utf-8')

def brief(result):
    o=result['order'];b=result['recommendation'];cat=result['catalog']
    lines=['# 运输方案比价与优化简报','',f"订单编号：{o['id']}",f"路线：{o['origin']} → {o['destination']}",f"重量：{o['weight']} 报价单位",f"时效要求：{str(o['max_days'])+' 天以内' if o['max_days'] is not None else '不限'}",'',f"报价文件：{cat['name']}；SHA-256：{cat['sha256']}",'',f"匹配 {result['matched']} 个方案，满足时效 {result['eligible']} 个。",'']
    if b:
        lines.extend([f"推荐：{b['carrier']} / {b['service']} / {b['mode']}",f"总成本：{b['cost']} 货币单位；运输天数：{b['days']}",f"核算：{b['formula']}",f"报价证据：原表第 {b['source_row']} 行（包含表头）",f"相对最快匹配方案的费用差：{result['savings_vs_fastest']} 货币单位。此为模拟比较，不是实际企业节省。",''])
    else: lines.extend([result['message'],''])
    lines.extend(['|承运商|服务|方式|天数|成本|满足时效|原表行号|','|---|---|---|---:|---:|---|---:|'])
    for c in result['candidates']:lines.append(f"|{c['carrier']}|{c['service']}|{c['mode']}|{c['days']}|{c['cost']}|{'是' if c['eligible'] else '否'}|{c['source_row']}|")
    lines.extend(['','## 口径与验证范围','', '数据来源：比赛附件 FreightRates.csv；演示订单为人为构造的模拟订单。实际费率真实性及有效日期未经外部核验。','计费假设：总成本 = max(最低收费, 重量 × 单位费率)。最低收费不作为额外附加费叠加；报价未明确规则，需采购确认。','原表未标注重量单位、币种、燃油附加费、税费、体积重与容量。这里使用报价单位、货币单位，不自动推断公斤或人民币。','重量区间包含上下端点；保留原始 0 天运输报价，此值需要业务核验。按未舍入金额排序，显示四舍五入到两位小数。','相同费用按天数、承运商代码、原始行号排序。多订单为独立逐单最优，无跨单拼单、网络路线或容量优化。'])
    return '\n'.join(lines)

def html_report(result):
    o=result['order'];b=result['recommendation'];e=lambda x:html.escape(str(x))
    rows=''.join('<tr>'+''.join(f'<td>{e(v)}</td>' for v in [c['carrier'],c['service'],c['mode'],c['days'],c['cost'],'推荐' if c['recommended'] else c['reason'],c['source_row']])+'</tr>' for c in result['candidates'])
    paragraphs=''.join((f'<h1>{e(line[2:])}</h1>' if line.startswith('# ') else f'<h2>{e(line[3:])}</h2>' if line.startswith('## ') else f'<p>{e(line)}</p>') for line in brief(result).splitlines() if line and not line.startswith('|'))
    return ('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>运输比价简报</title><style>body{max-width:1000px;margin:40px auto;padding:0 24px;font:16px/1.8 system-ui;color:#18343b}table{width:100%;border-collapse:collapse}td,th{padding:10px;border-bottom:1px solid #ddd;text-align:left}h1{color:#087c72}@media print{body{margin:0}}</style>'+paragraphs+'<table><thead><tr>'+''.join('<th>'+x+'</th>' for x in ['承运商','服务','方式','天数','费用','结论','原表行号'])+'</tr></thead><tbody>'+rows+'</tbody></table></html>').encode()

def result_zip(result):
    out=io.BytesIO()
    if 'results' in result:
        batch=result;rows=[[r['order']['id'],r['status'],r.get('recommendation',{}).get('carrier','') if r.get('recommendation') else '',r['recommendation']['cost'] if r.get('recommendation') else '',r['recommendation']['days'] if r.get('recommendation') else '',r.get('message',''),r['input_row']] for r in batch['results']]
        files={'批量推荐清单.csv':csv_bytes(['订单编号','状态','推荐承运商','费用_货币单位','天数','说明','输入行号'],rows),'批量简报.md':('# 批量运输比价简报\n\n'+json.dumps(batch['summary'],ensure_ascii=False,indent=2)+'\n\n原始报价 SHA-256：'+batch['catalog']['sha256']).encode()}
        for i,r in enumerate(batch['results']):
            if r['status']!='invalid':files[f'逐单简报/{i+1:03}.md']=brief(r).encode()
    else:
        fields=['承运商','服务等级','运输方式','天数','总成本_货币单位','原始未舍入成本','最低收费','单位费率','计费公式','满足时效','推荐','Pareto前沿','原表行号']
        rows=[[c['carrier'],c['service'],c['mode'],c['days'],c['cost'],c['raw_cost'],c['minimum_cost'],c['unit_rate'],c['formula'],c['eligible'],c['recommended'],c['pareto'],c['source_row']] for c in result['candidates']]
        files={'运输方案简报.md':brief(result).encode(),'运输方案简报.html':html_report(result),'比价明细.csv':csv_bytes(fields,rows)}
    files['运行快照.json']=json.dumps(result,ensure_ascii=False,indent=2).encode()
    try:
        import openpyxl
        wb=openpyxl.Workbook();ws=wb.active;ws.title='比价明细'
        header=fields if 'results' not in result else ['订单编号','状态','推荐承运商','费用_货币单位','天数','说明','输入行号']
        ws.append(header)
        numeric={'天数','原表行号','输入行号','总成本_货币单位','原始未舍入成本','最低收费','单位费率','费用_货币单位'}
        from decimal import Decimal
        for row in rows:
            ws.append([Decimal(str(v)) if h in numeric and str(v) not in ['', 'None'] else safe_cell(v) for h,v in zip(header,row)])
        for row in ws.iter_rows(min_row=2):
            for h,cell in zip(header,row):
                if h in {'总成本_货币单位','费用_货币单位'}:cell.number_format='0.00'
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        from openpyxl.styles import Font,PatternFill
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='087C72')
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(65,max(14,max(len(str(c.value or '')) for c in col)+3))
        b=io.BytesIO();wb.save(b);files['比价结果.xlsx']=b.getvalue()
    except ImportError: pass
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
        for name,data in files.items():z.writestr(name,data)
    return out.getvalue()
